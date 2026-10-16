import importlib.util
import json
import os
import shlex
import subprocess
import sys
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from smoke_proxy_support import Fail
from verify_office_processes import OwnedProcesses

SCRIPT = Path(__file__).with_name('verify-office.py')
SPEC = importlib.util.spec_from_file_location('verify_office_cli', SCRIPT)
OFFICE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(OFFICE)


def test_document_command_rejects_flattened_workbook_with_formula_diagnostic(tmp_path):
    path = tmp_path / 'flattened.xlsx'
    workbook = Workbook()
    workbook.active.title = 'Main'
    workbook.active['A2'] = 4
    workbook.active['B2'] = 12
    workbook.create_sheet('Second')['A2'] = 15
    workbook.save(path)
    result = subprocess.run(
        [sys.executable, str(SCRIPT), '--check-document', str(path)], capture_output=True, text=True, timeout=30
    )
    assert result.returncode == 1
    assert 'XLSX formulas' in result.stderr
    assert 'PASS' not in result.stdout


def test_missing_libreoffice_fails_preflight_without_partial_pass(tmp_path):
    env = dict(
        os.environ,
        OFFICE_SOFFICE=str(tmp_path / 'absent-libreoffice'),
        OFFICE_BROKER_PYTHON=sys.executable,
        OFFICE_DOCKER=sys.executable,
        OFFICE_NODE=sys.executable,
    )
    result = subprocess.run(
        [sys.executable, str(SCRIPT), '--preflight'], env=env, capture_output=True, text=True, timeout=30
    )
    assert result.returncode == 1
    assert 'LibreOffice' in result.stderr
    assert 'PASS' not in result.stdout
    assert 'skip' not in result.stdout.lower()


@pytest.mark.parametrize(
    ('architecture', 'expected'),
    [
        ('x86_64', 'linux/amd64'),
        ('amd64', 'linux/amd64'),
        ('aarch64', 'linux/arm64'),
        ('arm64', 'linux/arm64'),
    ],
)
def test_index_platform_selection_matches_docker_engine_architecture(architecture, expected):
    assert OFFICE.select_platform({'os': 'linux', 'architecture': architecture}) == expected


def test_unsupported_engine_platform_fails_without_an_image_substitution():
    with pytest.raises(Fail, match='unsupported Docker engine'):
        OFFICE.select_platform({'os': 'windows', 'architecture': 'amd64'})


def test_preflight_child_does_not_load_synthetic_dotenv(tmp_path, monkeypatch):
    monkeypatch.delenv('PYTHON_DOTENV_DISABLED', raising=False)
    monkeypatch.delenv('OCU_OFFICE_SYNTHETIC_DOTENV', raising=False)
    fixture = tmp_path / '.env'
    fixture.write_text('OCU_OFFICE_SYNTHETIC_DOTENV=loaded\n')
    code = (
        'import json,os,sys; from dotenv import load_dotenv; loaded=load_dotenv(sys.argv[1]); '
        'sys.stdout.write(json.dumps({"dotenv_loaded":loaded, '
        '"sentinel":os.environ.get("OCU_OFFICE_SYNTHETIC_DOTENV","absent")}))'
    )
    result = OFFICE.probe([sys.executable, '-c', code, str(fixture)], 'synthetic dotenv boundary')
    assert json.loads(result) == {'dotenv_loaded': False, 'sentinel': 'absent'}


def _generate_command(destination):
    return subprocess.run(
        [
            'make',
            '--no-print-directory',
            'verify-office',
            'OFFICE_ARGS=--generate-samples ' + shlex.quote(str(destination)),
        ],
        cwd=SCRIPT.parent.parent,
        env=dict(os.environ, OFFICE_PYTHON=sys.executable),
        capture_output=True,
        text=True,
        timeout=30,
    )


@pytest.mark.parametrize('destination_kind', ['directory', 'file', 'directory-symlink', 'dangling-symlink'])
def test_sample_command_rejects_existing_destination_without_changing_entries(tmp_path, destination_kind):
    owned = tmp_path / 'owned'
    owned.mkdir()
    (owned / 'sample.docx').write_bytes(b'Existing DOCX sentinel')
    (owned / 'sample.xlsx').write_bytes(b'Existing XLSX sentinel')
    (owned / 'nested').mkdir()
    (owned / 'nested' / 'unrelated.txt').write_bytes(b'Unrelated nested sentinel')
    missing = tmp_path / 'absent-target'
    destination = tmp_path / 'destination'
    if destination_kind == 'directory':
        destination = owned
    elif destination_kind == 'file':
        destination.write_bytes(b'Existing output-path file')
    else:
        target = owned if destination_kind == 'directory-symlink' else missing
        destination.symlink_to(target, target_is_directory=True)

    result = _generate_command(destination)

    assert result.returncode != 0
    assert 'PASS' not in result.stdout
    assert sorted(str(path.relative_to(owned)) for path in owned.rglob('*')) == [
        'nested',
        'nested/unrelated.txt',
        'sample.docx',
        'sample.xlsx',
    ]
    assert (owned / 'sample.docx').read_bytes() == b'Existing DOCX sentinel'
    assert (owned / 'sample.xlsx').read_bytes() == b'Existing XLSX sentinel'
    assert (owned / 'nested' / 'unrelated.txt').read_bytes() == b'Unrelated nested sentinel'
    if destination_kind == 'file':
        assert destination.read_bytes() == b'Existing output-path file'
    elif destination_kind.endswith('symlink'):
        assert destination.is_symlink()
        assert destination.readlink() == (owned if destination_kind == 'directory-symlink' else missing)
    assert not missing.exists()


def test_sample_command_generates_openable_three_format_corpus_in_new_directory(tmp_path):
    destination = tmp_path / 'new samples'
    result = _generate_command(destination)
    assert result.returncode == 0, result.stderr
    paths = json.loads(result.stdout.splitlines()[-1])
    assert set(paths) == {'docx', 'xlsx', 'pptx'}
    assert sorted(path.name for path in destination.iterdir()) == ['sample.docx', 'sample.pptx', 'sample.xlsx']
    document = Document(paths['docx'])
    assert document.paragraphs[0].text == 'Editable introduction'
    assert document.paragraphs[3].text == '中文内容保持完整。'
    assert document.tables[0].cell(1, 0).text == '中文数据'
    workbook = load_workbook(paths['xlsx'], data_only=False)
    try:
        assert workbook.sheetnames == ['Main', 'Second']
        assert workbook['Main']['A2'].value == 2
        assert workbook['Main']['B2'].value == '=A2*3'
        assert workbook['Second']['A2'].value == '=Main!B2+3'
    finally:
        workbook.close()
    assert [slide.shapes.title.text for slide in Presentation(paths['pptx']).slides] == [
        'Slide One',
        'Slide Two',
        'Slide Three',
    ]


def test_browser_roundtrip_preserves_primary_docx_failure_with_teardown_errors(tmp_path, monkeypatch):
    primary = 'DOCX open: real browser operation failed'
    private_log = (
        'DOCX open: real browser operation failed; unvalidated extra detail\n'
        + primary
        + '\nError: private editor URL https://document-server.invalid/?token=synthetic-private-token\n'
        'browser teardown: real browser operation failed\n'
        'DOCX native editor diagnostics: real browser operation failed\n'
    )
    executable = tmp_path / 'browser-boundary'
    executable.write_text(
        f'#!{sys.executable}\n'
        'import pathlib, sys, time\n'
        'profiles = [arg.split("=", 1)[1] for arg in sys.argv if arg.startswith("--user-data-dir=")]\n'
        'if profiles:\n'
        '    (pathlib.Path(profiles[0]) / "DevToolsActivePort").write_text('
        '"12345\\n/devtools/browser/01234567-89ab-cdef-0123-456789abcdef\\n")\n'
        '    time.sleep(60)\n'
        'else:\n'
        f'    sys.stderr.write({private_log!r})\n'
        '    sys.exit(1)\n'
    )
    executable.chmod(0o700)
    run = OFFICE.OfficeRun({'CHROMIUM': str(executable), 'OFFICE_NODE': str(executable)})
    run.scratch = tmp_path
    run.processes = OwnedProcesses(tmp_path)
    run.origin = 'http://broker.invalid'
    run.ds_origin = 'http://document-server.invalid'
    run.token = 'synthetic-private-token'
    run.chat = 'synthetic-chat'
    run.samples = {}
    monkeypatch.setattr(OFFICE, 'urlopen', lambda request, timeout: BytesIO(b'{"files": []}'))

    try:
        with pytest.raises(Fail) as failure:
            run.browser_roundtrip()

        assert str(failure.value) == primary
        browser_logs = list(tmp_path.glob('*-browser.log'))
        assert len(browser_logs) == 1
        assert browser_logs[0].read_text() == private_log
    finally:
        run.processes.close()


@pytest.fixture
def qualified_source(tmp_path):
    checkout = tmp_path / 'ocu'
    recipe = checkout / 'deploy/documentserver'
    recipe.mkdir(parents=True)
    (recipe / 'Dockerfile').write_text('FROM immutable-qualified-base\n')

    def git(*args):
        return subprocess.run(
            ['git', '-C', str(checkout), *args],
            check=True,
            capture_output=True,
            text=True,
        ).stdout.strip()

    git('init')
    git('config', 'user.name', 'Qualification fixture')
    git('config', 'user.email', 'qualification@example.invalid')
    git('add', '.')
    git('commit', '-m', 'recipe')
    tree = git('rev-parse', 'HEAD:deploy/documentserver')
    config = 'sha256:' + '1' * 64
    manifest = 'sha256:' + '2' * 64
    declaration = {
        'schema_version': 1,
        'recipe_tree': tree,
        'platforms': {
            'linux/arm64': {
                'reference': 'ocu-documentserver:' + '1' * 64,
                'configuration_digest': config,
                'platform_manifest_digest': manifest,
            }
        },
    }
    declaration_path = checkout / OFFICE.QUALIFIED_IMAGE_PATH

    def commit(payload=declaration):
        if payload is None:
            declaration_path.unlink(missing_ok=True)
        else:
            declaration_path.write_text(json.dumps(payload))
        git('add', '.')
        git('commit', '--allow-empty', '-m', 'qualification')
        return {'checkout': checkout, 'sha': git('rev-parse', 'HEAD')}

    source = commit()
    executable = tmp_path / 'docker-boundary'
    calls = tmp_path / 'calls.jsonl'
    state = tmp_path / 'image.json'
    state.write_text(
        json.dumps(
            {
                'Id': config,
                'Os': 'linux',
                'Architecture': 'arm64',
                'RepoDigests': [],
            }
        )
    )
    executable.write_text(
        f'#!{sys.executable}\n'
        'import json, pathlib, sys\n'
        f'with open({str(calls)!r}, "a") as output:\n'
        '    output.write(json.dumps(sys.argv[1:]) + "\\n")\n'
        'if sys.argv[1] == "info":\n'
        '    print(json.dumps({"os": "linux", "architecture": "aarch64"}))\n'
        'elif sys.argv[1:3] == ["image", "inspect"]:\n'
        f'    image = json.loads(pathlib.Path({str(state)!r}).read_text())\n'
        '    if image is None:\n'
        '        sys.stderr.write("image missing locally\\n")\n'
        '        sys.exit(1)\n'
        '    print(json.dumps([image]))\n'
        'else:\n'
        '    sys.stderr.write("unexpected Docker mutation\\n")\n'
        '    sys.exit(9)\n'
    )
    executable.chmod(0o700)
    return {
        'source': source,
        'commit': commit,
        'declaration': declaration,
        'declaration_path': declaration_path,
        'recipe': recipe,
        'tools': {'OFFICE_DOCKER': str(executable)},
        'calls': calls,
        'state': state,
        'config': config,
        'manifest': manifest,
    }


@pytest.mark.parametrize('identity_kind', ['config', 'manifest'])
def test_local_qualified_image_accepts_typed_identity_without_registry_digest(qualified_source, identity_kind):
    fixture = qualified_source
    image = json.loads(fixture['state'].read_text())
    image['Id'] = fixture[identity_kind]
    fixture['state'].write_text(json.dumps(image))
    # Working-tree qualification/recipe changes cannot override the pinned commit.
    fixture['declaration_path'].write_text('{}')
    (fixture['recipe'] / 'Dockerfile').write_text('FROM dirty:unqualified\n')
    inputs = OFFICE.image_preflight(fixture['tools'], fixture['source'])
    assert inputs['image_identity'] == fixture[identity_kind]
    assert inputs['image_reference'] == 'ocu-documentserver:' + '1' * 64
    assert inputs['platform'] == 'linux/arm64'
    assert inputs['image_qualification']['configuration_digest'] == fixture['config']
    calls = [json.loads(line) for line in fixture['calls'].read_text().splitlines()]
    assert calls[1] == [
        'image',
        'inspect',
        '--platform',
        'linux/arm64',
        'ocu-documentserver:' + '1' * 64,
    ]
    assert [call[0] for call in calls] == ['info', 'image']


def test_qualified_configuration_only_supports_classic_image_store(qualified_source):
    fixture = qualified_source
    del fixture['declaration']['platforms']['linux/arm64']['platform_manifest_digest']
    source = fixture['commit']()
    inputs = OFFICE.image_preflight(fixture['tools'], source)
    assert inputs['image_identity'] == fixture['config']


def _apply_qualification_source_defect(fixture, defect):
    declaration = fixture['declaration']
    record = declaration['platforms']['linux/arm64']
    if defect == 'missing-declaration':
        declaration = None
    elif defect == 'stale-recipe':
        (fixture['recipe'] / 'Dockerfile').write_text('FROM changed:recipe\n')
    elif defect == 'schema':
        declaration['schema_version'] = True
    elif defect == 'unsupported-platform':
        declaration['platforms']['linux/riscv64'] = dict(record)
    elif defect == 'missing-platform':
        declaration['platforms'] = {'linux/amd64': record}
    elif defect == 'mutable-reference':
        record['reference'] = 'ocu-documentserver:latest'
    elif defect == 'bad-config':
        record['configuration_digest'] = 'not-a-digest'
    elif defect == 'bad-manifest':
        record['platform_manifest_digest'] = 'sha256:short'
    elif defect == 'index-field':
        record['index_digest'] = 'sha256:' + '3' * 64
    return declaration


@pytest.mark.parametrize(
    'defect',
    [
        'missing-declaration',
        'stale-recipe',
        'schema',
        'unsupported-platform',
        'missing-platform',
        'mutable-reference',
        'bad-config',
        'bad-manifest',
        'index-field',
        'missing-image',
        'wrong-platform',
        'rebound-id',
        'index-id',
    ],
)
def test_unqualified_source_or_image_refuses_before_runtime_resources(qualified_source, defect):
    fixture = qualified_source
    declaration = _apply_qualification_source_defect(fixture, defect)
    image = json.loads(fixture['state'].read_text())
    if defect == 'missing-image':
        image = None
    elif defect == 'wrong-platform':
        image['Architecture'] = 'amd64'
    elif defect == 'rebound-id':
        image['Id'] = 'sha256:' + '9' * 64
    elif defect == 'index-id':
        image['Id'] = 'sha256:' + '3' * 64
        image['RepoDigests'] = ['onlyoffice/documentserver@sha256:' + '3' * 64]
    source = fixture['commit'](declaration)
    fixture['state'].write_text(json.dumps(image))
    with pytest.raises(Fail):
        OFFICE.image_preflight(fixture['tools'], source)
    calls = (
        [json.loads(line) for line in fixture['calls'].read_text().splitlines()] if fixture['calls'].exists() else []
    )
    assert all(call[0] == 'info' or call[:2] == ['image', 'inspect'] for call in calls)
    if defect not in {'missing-platform', 'missing-image', 'wrong-platform', 'rebound-id', 'index-id'}:
        assert calls == []
