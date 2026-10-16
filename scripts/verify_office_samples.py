"""Synthetic Office documents; ZIP and document metadata are fixed across runs."""

import datetime as dt
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.shared import Inches as DocxInches
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from smoke_proxy_support import Fail
from verify_office_fonts import FONT, apply_font_policy

STAMP = dt.datetime(2000, 1, 1, tzinfo=dt.UTC)


def _picture() -> BytesIO:
    image = Image.new('RGB', (16, 16), '#2364aa')
    for x in range(8):
        for y in range(16):
            image.putpixel((x, y), (220, 40, 50))
    stream = BytesIO()
    image.save(stream, format='PNG')
    stream.seek(0)
    return stream


def _package(document, path: Path, properties) -> None:
    properties.created = STAMP
    for prefix, namespace in (
        ('cp', 'http://schemas.openxmlformats.org/package/2006/metadata/core-properties'),
        ('dc', 'http://purl.org/dc/elements/1.1/'),
        ('dcterms', 'http://purl.org/dc/terms/'),
        ('xsi', 'http://www.w3.org/2001/XMLSchema-instance'),
    ):
        ElementTree.register_namespace(prefix, namespace)
    properties.modified = STAMP
    stream = BytesIO()
    document.save(stream)
    with ZipFile(BytesIO(stream.getvalue())) as source, ZipFile(path, 'w') as output:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            content = source.read(name)
            if name == 'docProps/core.xml':
                core = ElementTree.fromstring(content)
                for field in ('created', 'modified'):
                    node = core.find(f'{{http://purl.org/dc/terms/}}{field}')
                    if node is not None:
                        node.text = '2000-01-01T00:00:00Z'
                content = ElementTree.tostring(core, encoding='utf-8', xml_declaration=True)
            if name.startswith(('word/', 'xl/', 'ppt/')) and name.endswith('.xml'):
                content = apply_font_policy(content)
            output.writestr(info, content)


def generate_samples(directory: Path) -> dict[str, Path]:
    try:
        directory.mkdir(mode=0o700, parents=True, exist_ok=False)
    except FileExistsError:
        raise Fail('sample generation: output directory already exists; choose a new destination') from None
    paths = {kind: directory / f'sample.{kind}' for kind in ('docx', 'xlsx', 'pptx')}
    doc = Document()
    doc.add_paragraph('Editable introduction')
    doc.add_heading('Office verification', level=1)
    doc.add_heading('Content preservation', level=2)
    doc.add_paragraph('中文内容保持完整。')
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (('Item', 'Count'), ('中文数据', '3')), strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    doc.add_picture(_picture(), width=DocxInches(0.5))
    doc.add_paragraph('OCU_FONT_PROBE_DOCX').runs[0].font.name = FONT
    _package(doc, paths['docx'], doc.core_properties)

    book = Workbook()
    main = book.active
    main.title = 'Main'
    main.append(['Input', 'Triple'])
    main.append([2, '=A2*3'])
    main['A5'] = 'OCU_FONT_PROBE_XLSX'
    main['A5'].font = Font(name=FONT)
    main.column_dimensions['A'].width = 32
    second = book.create_sheet('Second')
    second.append(['Cross-sheet total'])
    second.append(['=Main!B2+3'])
    chart = BarChart()
    chart.add_data(Reference(main, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    main.add_chart(chart, 'D2')
    _package(book, paths['xlsx'], book.properties)
    book.close()

    deck = Presentation()
    for title in ('Slide One', 'Slide Two', 'Slide Three'):
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = title
    shape = deck.slides[0].shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(5.5), Inches(3), Inches(0.6))
    shape.text = 'Preserved shape'
    deck.slides[1].shapes.add_picture(_picture(), Inches(1), Inches(5), width=Inches(0.5))
    probe = deck.slides[2].shapes.add_textbox(Inches(1), Inches(5), Inches(7), Inches(0.6))
    probe.text = 'OCU_FONT_PROBE_PPTX'
    probe.text_frame.paragraphs[0].runs[0].font.name = FONT
    _package(deck, paths['pptx'], deck.core_properties)
    return paths
