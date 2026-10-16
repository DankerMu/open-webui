from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from smoke_proxy_support import Fail
from verify_office_fonts import declared_font_policy
from verify_office_samples import generate_samples

POLICY = 'OCURequiredMissingFont172'


def test_generated_corpus_has_uniform_declared_and_effective_text_fonts(tmp_path):
    samples = generate_samples(tmp_path / 'samples')
    assert {kind: declared_font_policy(path)['font'] for kind, path in samples.items()} == {
        'docx': POLICY,
        'xlsx': POLICY,
        'pptx': POLICY,
    }
    document = Document(samples['docx'])
    assert document.paragraphs[0].runs[0].font.name == POLICY
    assert document.paragraphs[1].runs[0].font.name == POLICY
    assert document.paragraphs[3].runs[0].font.name == POLICY
    assert document.tables[0].cell(1, 0).paragraphs[0].runs[0].font.name == POLICY
    workbook = load_workbook(samples['xlsx'])
    try:
        assert workbook['Main']['A1'].font.name == POLICY
        assert workbook['Second']['A2'].font.name == POLICY
    finally:
        workbook.close()
    slides = Presentation(samples['pptx']).slides
    assert [slide.shapes.title.text_frame.paragraphs[0].runs[0].font.name for slide in slides] == [POLICY] * 3


def test_mixed_published_declarations_are_an_attribution_blocker_not_invented_substitutions(tmp_path):
    original = generate_samples(tmp_path / 'input')['docx']
    published = tmp_path / 'mixed.docx'
    with ZipFile(original) as source, ZipFile(published, 'w') as output:
        for info in source.infolist():
            content = source.read(info)
            if info.filename == 'word/document.xml':
                content = content.replace(POLICY.encode(), b'Courier', 1)
            output.writestr(info, content)
    with pytest.raises(Fail, match='fidelity attribution.*mixed policies'):
        declared_font_policy(published)


def test_missing_word_run_and_default_fonts_block_attribution_even_with_uniform_theme(tmp_path):
    original = generate_samples(tmp_path / 'input')['docx']
    damaged = tmp_path / 'missing-defaults.docx'
    namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    with ZipFile(original) as source, ZipFile(damaged, 'w') as output:
        for info in source.infolist():
            content = source.read(info)
            if info.filename in ('word/document.xml', 'word/styles.xml'):
                root = etree.fromstring(content)
                if info.filename == 'word/document.xml':
                    declarations = list(root.iter(namespace + 'rFonts'))
                else:
                    declarations = root.findall(
                        './/'
                        + namespace
                        + 'docDefaults/'
                        + namespace
                        + 'rPrDefault/'
                        + namespace
                        + 'rPr/'
                        + namespace
                        + 'rFonts',
                    )
                for declaration in declarations:
                    declaration.getparent().remove(declaration)
                content = etree.tostring(root)
            output.writestr(info, content)
    with pytest.raises(Fail, match='unresolved Word text/default fonts'):
        declared_font_policy(damaged)


def test_missing_theme_slot_cannot_be_treated_as_a_resolved_chart_font(tmp_path):
    original = generate_samples(tmp_path / 'input')['xlsx']
    damaged = tmp_path / 'missing-theme-slot.xlsx'
    namespace = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    with ZipFile(original) as source, ZipFile(damaged, 'w') as output:
        for info in source.infolist():
            content = source.read(info)
            if info.filename == 'xl/theme/theme1.xml':
                root = etree.fromstring(content)
                declaration = root.find('.//' + namespace + 'majorFont/' + namespace + 'latin')
                declaration.getparent().remove(declaration)
                content = etree.tostring(root)
            output.writestr(info, content)
    with pytest.raises(Fail, match='unresolved majorFont/latin font slot'):
        declared_font_policy(damaged)


def test_unused_conflicting_word_style_does_not_poison_explicit_text_policy(tmp_path):
    original = generate_samples(tmp_path / 'input')['docx']
    published = tmp_path / 'native-unused-table-style.docx'
    namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    with ZipFile(original) as source, ZipFile(published, 'w') as output:
        for info in source.infolist():
            content = source.read(info)
            if info.filename == 'word/styles.xml':
                root = etree.fromstring(content)
                style = etree.SubElement(root, namespace + 'style')
                style.set(namespace + 'type', 'table')
                style.set(namespace + 'styleId', 'UnselectedNativeTableStyle')
                conditional = etree.SubElement(style, namespace + 'tblStylePr')
                conditional.set(namespace + 'type', 'firstRow')
                properties = etree.SubElement(conditional, namespace + 'rPr')
                declaration = etree.SubElement(properties, namespace + 'rFonts')
                for slot in ('ascii', 'hAnsi'):
                    declaration.set(namespace + slot, 'UnusedTableFamily')
                content = etree.tostring(root)
            output.writestr(info, content)
    assert declared_font_policy(published)['font'] == POLICY


@pytest.mark.parametrize('orphan_kind', ['conflicting', 'missing-slot', 'malformed-xml'])
def test_unreferenced_spreadsheet_theme_does_not_poison_selected_theme(tmp_path, orphan_kind):
    original = generate_samples(tmp_path / 'input')['xlsx']
    published = tmp_path / 'native-orphan-theme.xlsx'
    namespace = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    root = etree.Element(namespace + 'theme')
    elements = etree.SubElement(root, namespace + 'themeElements')
    scheme = etree.SubElement(elements, namespace + 'fontScheme')
    for role in ('majorFont', 'minorFont'):
        group = etree.SubElement(scheme, namespace + role)
        for slot in ('latin', 'ea', 'cs'):
            declaration = etree.SubElement(group, namespace + slot)
            declaration.set('typeface', '' if orphan_kind == 'missing-slot' and slot == 'ea' else 'UnusedThemeFamily')
    orphan = b'<malformed orphan theme' if orphan_kind == 'malformed-xml' else etree.tostring(root)
    with ZipFile(original) as source, ZipFile(published, 'w') as output:
        for info in source.infolist():
            output.writestr(info, source.read(info))
        output.writestr('xl/theme/theme.xml', orphan)
    assert declared_font_policy(published)['font'] == POLICY


def _rewrite_package(original, published, replacements):
    with ZipFile(original) as source, ZipFile(published, 'w') as output:
        names = set(source.namelist())
        for info in source.infolist():
            output.writestr(info, replacements.get(info.filename, source.read(info)))
        for name, content in replacements.items():
            if name not in names:
                output.writestr(name, content)


def _word_font_declaration(properties, namespace, family):
    fonts = etree.SubElement(properties, namespace + 'rFonts')
    for slot in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(namespace + slot, family)


def _inherited_word_package(tmp_path, override=None, direct=False):
    original = generate_samples(tmp_path / 'input')['docx']
    published = tmp_path / 'inherited-word.docx'
    namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    document = etree.Element(namespace + 'document')
    body = etree.SubElement(document, namespace + 'body')
    parent = body
    if override in ('table', 'default-table'):
        table = etree.SubElement(body, namespace + 'tbl')
        properties = etree.SubElement(table, namespace + 'tblPr')
        etree.SubElement(properties, namespace + 'tblLook').set(namespace + 'firstRow', '1')
        if override == 'table':
            etree.SubElement(properties, namespace + 'tblStyle').set(namespace + 'val', 'ActiveAlternate')
        parent = etree.SubElement(etree.SubElement(table, namespace + 'tr'), namespace + 'tc')
    paragraph = etree.SubElement(parent, namespace + 'p')
    run = etree.SubElement(paragraph, namespace + 'r')
    etree.SubElement(run, namespace + 't').text = 'Effective inherited text'
    styles = etree.Element(namespace + 'styles')
    defaults = etree.SubElement(styles, namespace + 'docDefaults')
    default_run = etree.SubElement(etree.SubElement(defaults, namespace + 'rPrDefault'), namespace + 'rPr')
    _word_font_declaration(default_run, namespace, POLICY)
    style = etree.SubElement(styles, namespace + 'style')
    style.set(
        namespace + 'type',
        {'character': 'character', 'table': 'table', 'default-table': 'table'}.get(override, 'paragraph'),
    )
    style.set(namespace + 'styleId', 'ActiveAlternate')
    if override in ('default-paragraph', 'default-table'):
        style.set(namespace + 'default', '1')
    if override in ('table', 'default-table'):
        style = etree.SubElement(style, namespace + 'tblStylePr')
        style.set(namespace + 'type', 'firstRow')
    _word_font_declaration(etree.SubElement(style, namespace + 'rPr'), namespace, 'ActiveAlternateFamily')
    if override == 'character':
        properties = etree.SubElement(run, namespace + 'rPr')
        etree.SubElement(properties, namespace + 'rStyle').set(namespace + 'val', 'ActiveAlternate')
    elif override == 'paragraph':
        properties = etree.SubElement(paragraph, namespace + 'pPr')
        etree.SubElement(properties, namespace + 'pStyle').set(namespace + 'val', 'ActiveAlternate')
    elif override == 'numbering':
        properties = etree.SubElement(paragraph, namespace + 'pPr')
        numbering = etree.SubElement(properties, namespace + 'numPr')
        etree.SubElement(numbering, namespace + 'numId').set(namespace + 'val', '1')
    if direct:
        properties = etree.SubElement(run, namespace + 'rPr')
        _word_font_declaration(properties, namespace, POLICY)
    _rewrite_package(
        original,
        published,
        {
            'word/document.xml': etree.tostring(document),
            'word/styles.xml': etree.tostring(styles),
        },
    )
    return published


def test_document_defaults_prove_word_text_only_without_applicable_font_overrides(tmp_path):
    published = _inherited_word_package(tmp_path)
    assert declared_font_policy(published)['font'] == POLICY


@pytest.mark.parametrize(
    'override', ['character', 'paragraph', 'default-paragraph', 'table', 'default-table', 'numbering']
)
def test_active_word_inheritance_cannot_borrow_unrelated_uniform_defaults(tmp_path, override):
    published = _inherited_word_package(tmp_path, override)
    with pytest.raises(Fail, match='unresolved Word text/default fonts'):
        declared_font_policy(published)


@pytest.mark.parametrize('override', ['table', 'default-table'])
def test_complete_direct_fonts_outrank_selected_conditional_table_font(tmp_path, override):
    published = _inherited_word_package(tmp_path, override, direct=True)
    assert declared_font_policy(published)['font'] == POLICY


@pytest.mark.parametrize('target', ['./theme/../theme/selected-native.xml', '/xl/theme/selected-native.xml'])
def test_workbook_relationship_selects_theme_not_conventional_part_name(tmp_path, target):
    original = generate_samples(tmp_path / 'input')['xlsx']
    published = tmp_path / 'selected-native-theme.xlsx'
    namespace = '{http://schemas.openxmlformats.org/package/2006/relationships}'
    with ZipFile(original) as package:
        relationships = etree.fromstring(package.read('xl/_rels/workbook.xml.rels'))
        theme = next(
            row for row in relationships.findall(namespace + 'Relationship') if row.get('Type').endswith('/theme')
        )
        theme.set('Target', target)
        selected = package.read('xl/theme/theme1.xml')
    _rewrite_package(
        original,
        published,
        {
            'xl/_rels/workbook.xml.rels': etree.tostring(relationships),
            'xl/theme/theme1.xml': b'<malformed unselected conventional theme',
            'xl/theme/selected-native.xml': selected,
        },
    )
    assert declared_font_policy(published)['font'] == POLICY


def test_reachable_alternate_theme_remains_a_mixed_font_attribution_failure(tmp_path):
    original = generate_samples(tmp_path / 'input')['xlsx']
    published = tmp_path / 'selected-alternate-theme.xlsx'
    namespace = '{http://schemas.openxmlformats.org/package/2006/relationships}'
    with ZipFile(original) as package:
        relationships = etree.fromstring(package.read('xl/_rels/workbook.xml.rels'))
        theme = next(
            row for row in relationships.findall(namespace + 'Relationship') if row.get('Type').endswith('/theme')
        )
        theme.set('Target', 'theme/active-alternate.xml')
        selected = package.read('xl/theme/theme1.xml').replace(POLICY.encode(), b'ActiveAlternateFamily')
    _rewrite_package(
        original,
        published,
        {
            'xl/_rels/workbook.xml.rels': etree.tostring(relationships),
            'xl/theme/active-alternate.xml': selected,
        },
    )
    with pytest.raises(Fail, match='fidelity attribution.*mixed policies'):
        declared_font_policy(published)


@pytest.mark.parametrize('fault', ['missing', 'escaping', 'encoded-escaping', 'backslash', 'duplicate-id', 'external'])
def test_selected_theme_relationship_requires_unambiguous_internal_material(tmp_path, fault):
    original = generate_samples(tmp_path / 'input')['xlsx']
    published = tmp_path / 'unproven-selected-theme.xlsx'
    namespace = '{http://schemas.openxmlformats.org/package/2006/relationships}'
    with ZipFile(original) as package:
        relationships = etree.fromstring(package.read('xl/_rels/workbook.xml.rels'))
    theme = next(row for row in relationships.findall(namespace + 'Relationship') if row.get('Type').endswith('/theme'))
    if fault == 'duplicate-id':
        duplicate = etree.SubElement(relationships, namespace + 'Relationship')
        duplicate.attrib.update(theme.attrib)
    elif fault == 'external':
        theme.set('TargetMode', 'External')
        theme.set('Target', 'https://unavailable.example.invalid/theme.xml')
    else:
        theme.set(
            'Target',
            {
                'missing': 'theme/absent.xml',
                'escaping': '../../outside.xml',
                'encoded-escaping': '%2e%2e/%2e%2e/outside.xml',
                'backslash': r'theme\outside.xml',
            }[fault],
        )
    _rewrite_package(original, published, {'xl/_rels/workbook.xml.rels': etree.tostring(relationships)})
    with pytest.raises(Fail, match='fidelity attribution:.*(missing|ambiguous|escaping|unsafe|external)'):
        declared_font_policy(published)


@pytest.mark.parametrize('fault', ['missing', 'ambiguous', 'external'])
def test_package_requires_one_internal_office_document_authority(tmp_path, fault):
    original = generate_samples(tmp_path / 'input')['xlsx']
    published = tmp_path / 'unproven-main-document.xlsx'
    namespace = '{http://schemas.openxmlformats.org/package/2006/relationships}'
    with ZipFile(original) as package:
        relationships = etree.fromstring(package.read('_rels/.rels'))
    main = next(
        row for row in relationships.findall(namespace + 'Relationship') if row.get('Type').endswith('/officeDocument')
    )
    if fault == 'missing':
        relationships.remove(main)
    elif fault == 'ambiguous':
        duplicate = etree.SubElement(relationships, namespace + 'Relationship')
        duplicate.attrib.update(main.attrib)
        duplicate.set('Id', 'another-office-document')
    else:
        main.set('TargetMode', 'External')
        main.set('Target', 'https://unavailable.example.invalid/workbook.xml')
    _rewrite_package(original, published, {'_rels/.rels': etree.tostring(relationships)})
    with pytest.raises(Fail, match='fidelity attribution:.*officeDocument'):
        declared_font_policy(published)


def test_malformed_reachable_relationships_cannot_silently_drop_font_authority(tmp_path):
    original = generate_samples(tmp_path / 'input')['docx']
    published = tmp_path / 'malformed-relationship-authority.docx'
    _rewrite_package(original, published, {'word/_rels/document.xml.rels': b'<not-relationships/>'})
    with pytest.raises(Fail, match='fidelity attribution:.*malformed relationships'):
        declared_font_policy(published)


def _presentation_with_distinct_themes(tmp_path, conflicting=False):
    original = generate_samples(tmp_path / 'input')['pptx']
    published = tmp_path / 'native-two-themes.pptx'
    namespace = '{http://schemas.openxmlformats.org/package/2006/relationships}'
    theme_type = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme'
    with ZipFile(original) as package:
        relationships = etree.fromstring(package.read('ppt/_rels/presentation.xml.rels'))
        secondary = etree.fromstring(package.read('ppt/theme/theme1.xml'))
    for row in list(relationships):
        if row.get('Type') == theme_type:
            relationships.remove(row)
    for identifier, target in (
        ('native-theme-one', 'theme/theme1.xml'),
        ('native-theme-two', 'theme/native-secondary.xml'),
    ):
        etree.SubElement(
            relationships,
            namespace + 'Relationship',
            {
                'Id': identifier,
                'Type': theme_type,
                'Target': target,
            },
        )
    secondary.set('name', 'Native secondary theme')
    content = etree.tostring(secondary)
    if conflicting:
        content = content.replace(POLICY.encode(), b'ActiveAlternateFamily')
    _rewrite_package(
        original,
        published,
        {
            'ppt/_rels/presentation.xml.rels': etree.tostring(relationships),
            'ppt/theme/native-secondary.xml': content,
        },
    )
    return published


def test_distinct_reachable_presentation_themes_with_same_family_prove_uniform_policy(tmp_path):
    published = _presentation_with_distinct_themes(tmp_path)
    assert declared_font_policy(published)['font'] == POLICY


def test_conflicting_reachable_presentation_theme_cannot_claim_uniform_policy(tmp_path):
    published = _presentation_with_distinct_themes(tmp_path, conflicting=True)
    with pytest.raises(Fail):
        declared_font_policy(published)


def test_generated_numbering_label_fonts_are_not_overridden_by_direct_body_fonts(tmp_path):
    original = generate_samples(tmp_path / 'input')['docx']
    published = tmp_path / 'active-numbering-label.docx'
    namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    with ZipFile(original) as package:
        document = etree.fromstring(package.read('word/document.xml'))
    paragraph = next(document.iter(namespace + 'p'))
    properties = paragraph.find(namespace + 'pPr')
    if properties is None:
        properties = etree.Element(namespace + 'pPr')
        paragraph.insert(0, properties)
    selected = etree.SubElement(properties, namespace + 'numPr')
    etree.SubElement(selected, namespace + 'ilvl').set(namespace + 'val', '0')
    etree.SubElement(selected, namespace + 'numId').set(namespace + 'val', '1')
    numbering = etree.Element(namespace + 'numbering')
    abstract = etree.SubElement(numbering, namespace + 'abstractNum')
    abstract.set(namespace + 'abstractNumId', '0')
    level = etree.SubElement(abstract, namespace + 'lvl')
    level.set(namespace + 'ilvl', '0')
    for tag, value in (('start', '1'), ('numFmt', 'decimal'), ('lvlText', '%1.')):
        etree.SubElement(level, namespace + tag).set(namespace + 'val', value)
    _word_font_declaration(etree.SubElement(level, namespace + 'rPr'), namespace, 'GeneratedLabelAlternateFamily')
    instance = etree.SubElement(numbering, namespace + 'num')
    instance.set(namespace + 'numId', '1')
    etree.SubElement(instance, namespace + 'abstractNumId').set(namespace + 'val', '0')
    _rewrite_package(
        original,
        published,
        {
            'word/document.xml': etree.tostring(document),
            'word/numbering.xml': etree.tostring(numbering),
        },
    )
    with pytest.raises(Fail):
        declared_font_policy(published)
