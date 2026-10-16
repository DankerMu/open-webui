"""Read actual publications without repairing, recalculating or rewriting them."""

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

MC_NS = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
WPG_NS = 'http://schemas.microsoft.com/office/word/2010/wordprocessingGroup'


class DocumentCheckError(RuntimeError):
    def __init__(self, format: str, check: str, detail: str):
        self.format = format
        self.check = check
        super().__init__(f'{format} {check}: {detail}')


def require(condition, format, check, detail):
    if not condition:
        raise DocumentCheckError(format, check, detail)


def judge_xlsx_formulas(path: Path, *, expected_formulas, expected_values) -> dict:
    formulas = load_workbook(path, data_only=False, read_only=True)
    try:
        for (sheet, address), expected in expected_formulas.items():
            require(sheet in formulas.sheetnames, 'XLSX', 'sheets', f'missing {sheet}')
            cell = formulas[sheet][address]
            require(cell.data_type == 'f', 'XLSX', 'formulas', f'{sheet}!{address} is not a live formula')
            # Quoting a simple sheet name does not change the reference's meaning.
            actual = str(cell.value).replace("'Main'!", 'Main!')
            require(actual == expected, 'XLSX', 'formulas', f'{sheet}!{address} changed formula')
    finally:
        formulas.close()
    values = load_workbook(path, data_only=True, read_only=True)
    try:
        observed = {}
        for (sheet, address), expected in expected_values.items():
            require(sheet in values.sheetnames, 'XLSX', 'sheets', f'missing {sheet}')
            value = values[sheet][address].value
            require(
                type(value) in (int, float) and value == expected,
                'XLSX',
                'values',
                f'{sheet}!{address} expected {expected}, observed {value!r}',
            )
            observed[sheet, address] = value
        return observed
    finally:
        values.close()


def _judge_image(blob: bytes, format: str) -> None:
    with Image.open(BytesIO(blob)) as image:
        image = image.convert('RGB')
        require(image.size == (16, 16), format, 'image', 'image dimensions changed')
        pixels = image.get_flattened_data()
        require(
            all(pixel == ((220, 40, 50) if index % 16 < 8 else (35, 100, 170)) for index, pixel in enumerate(pixels)),
            format,
            'image',
            'image pixels changed',
        )


def judge_docx(path: Path) -> dict:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    headings = [
        (paragraph.style.name, paragraph.text)
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith('Heading')
    ]
    require(
        headings == [('Heading 1', 'Office verification'), ('Heading 2', 'Content preservation')],
        'DOCX',
        'headings',
        'required heading text or levels changed',
    )
    require(
        paragraphs[0] == 'OCU_DOCX_EDIT Editable introduction',
        'DOCX',
        'edit',
        'published introduction differs from the intended edit',
    )
    require('中文内容保持完整。' in paragraphs, 'DOCX', 'Chinese text', 'Chinese paragraph changed')
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    require(tables == [[['Item', 'Count'], ['中文数据', '3']]], 'DOCX', 'table', 'table content changed')
    inlines = [shape._inline for shape in document.inline_shapes]
    # DocumentServer wraps its ordinary DrawingML picture in a wpg-compatible choice.
    for alternate in document.element.body.iter(f'{{{MC_NS}}}AlternateContent'):
        if alternate.getparent().tag != qn('w:r'):
            continue
        for choice in alternate.findall(f'{{{MC_NS}}}Choice'):
            required = choice.get('Requires', '').split()
            if required and all(choice.nsmap.get(prefix) == WPG_NS for prefix in required):
                for drawing in choice.findall(qn('w:drawing')):
                    inlines.extend(drawing.findall(qn('wp:inline')))
                break
    require(len(inlines) == 1, 'DOCX', 'image', 'required inline image missing')
    try:
        picture = inlines[0].graphic.graphicData.pic
        require(picture is not None, 'DOCX', 'image', 'required inline picture missing')
        relationship = picture.blipFill.blip.embed
        _judge_image(document.part.related_parts[relationship].blob, 'DOCX')
    except (AttributeError, KeyError, OSError, ValueError) as exc:
        raise DocumentCheckError('DOCX', 'image', f'cannot read inline picture: {type(exc).__name__}') from exc
    return {'format': 'DOCX', 'checks': ['edit', 'headings', 'table', 'image', 'Chinese text']}


def judge_xlsx(path: Path) -> dict:
    values = judge_xlsx_formulas(
        path,
        expected_formulas={('Main', 'B2'): '=A2*3', ('Second', 'A2'): '=Main!B2+3'},
        expected_values={('Main', 'B2'): 12, ('Second', 'A2'): 15},
    )
    workbook = load_workbook(path, data_only=False)
    try:
        require(workbook.sheetnames == ['Main', 'Second'], 'XLSX', 'sheets', 'sheet names or order changed')
        require(workbook['Main']['A2'].value == 4, 'XLSX', 'edit', 'published input is not 4')
        require(len(workbook['Main']._charts) == 1, 'XLSX', 'chart', 'required chart missing')
    finally:
        workbook.close()
    return {
        'format': 'XLSX',
        'checks': ['edit', 'sheets', 'formulas', 'values', 'chart'],
        'values': {f'{sheet}!{cell}': value for (sheet, cell), value in values.items()},
    }


def judge_pptx(path: Path) -> dict:
    slides = Presentation(path).slides
    require(len(slides) == 3, 'PPTX', 'slide count', 'expected three slides')
    titles = [slide.shapes.title.text if slide.shapes.title is not None else None for slide in slides]
    require(titles[0] == 'Slide One OCU_PPTX_EDIT', 'PPTX', 'edit', 'edited first slide title missing')
    require(
        titles == ['Slide One OCU_PPTX_EDIT', 'Slide Two', 'Slide Three'],
        'PPTX',
        'slide order',
        'slide title order changed',
    )
    rectangles = [
        shape
        for shape in slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and shape.auto_shape_type == MSO_SHAPE.RECTANGLE
    ]
    require(
        len(rectangles) == 1 and rectangles[0].text == 'Preserved shape',
        'PPTX',
        'shapes',
        'required rectangle or text missing',
    )
    pictures = [shape for shape in slides[1].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    require(len(pictures) == 1, 'PPTX', 'image', 'required image missing from second slide')
    _judge_image(pictures[0].image.blob, 'PPTX')
    return {'format': 'PPTX', 'checks': ['edit', 'slide count', 'slide order', 'shapes', 'image']}


def judge_document(path: Path) -> dict:
    kind = path.suffix.lstrip('.').lower()
    judges = {'docx': judge_docx, 'xlsx': judge_xlsx, 'pptx': judge_pptx}
    require(kind in judges, kind.upper(), 'format', 'unsupported document format')
    try:
        return judges[kind](path)
    except DocumentCheckError:
        raise
    except Exception as exc:
        raise DocumentCheckError(kind.upper(), 'read', type(exc).__name__) from exc
