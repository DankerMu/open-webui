from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook
from pptx import Presentation
from verify_office_samples import generate_samples


def test_sample_bytes_and_metadata_are_stable_across_output_directories(tmp_path):
    first = generate_samples(tmp_path / 'first')
    second = generate_samples(tmp_path / 'second')
    for kind in ('docx', 'xlsx', 'pptx'):
        assert first[kind].read_bytes() == second[kind].read_bytes()
        with ZipFile(first[kind]) as package:
            core = ElementTree.fromstring(package.read('docProps/core.xml'))
            assert core.find('{http://purl.org/dc/terms/}modified').text == '2000-01-01T00:00:00Z'
    book = load_workbook(first['xlsx'], data_only=False)
    try:
        assert book.sheetnames == ['Main', 'Second']
        assert book['Main']['A2'].value == 2
        assert book['Main']['B2'].value == '=A2*3'
        assert book['Second']['A2'].value == '=Main!B2+3'
    finally:
        book.close()
    assert [slide.shapes.title.text for slide in Presentation(first['pptx']).slides] == [
        'Slide One',
        'Slide Two',
        'Slide Three',
    ]
