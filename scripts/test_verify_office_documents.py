from copy import deepcopy
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from verify_office_documents import (
    DocumentCheckError,
    judge_docx,
    judge_pptx,
    judge_xlsx,
    judge_xlsx_formulas,
)

FORMULAS = {('Main', 'B2'): '=A2*3', ('Second', 'A2'): '=Main!B2+3'}
EXPECTED_VALUES = {('Main', 'B2'): 12, ('Second', 'A2'): 15}
SHEET_NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
MC_NS = 'http://schemas.openxmlformats.org/markup-compatibility/2006'


def _published_workbook(path: Path, *, flattened: bool = False) -> None:
    workbook = Workbook()
    main = workbook.active
    main.title = 'Main'
    main.append(['Input', 'Triple'])
    main.append([4, '=A2*3'])
    second = workbook.create_sheet('Second')
    second.append(['Cross-sheet total'])
    second.append(['=Main!B2+3'])
    chart = BarChart()
    chart.add_data(Reference(main, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    main.add_chart(chart, 'D2')
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    # openpyxl does not calculate caches; these are the worked results of 4*3 and 12+3.
    caches = {'xl/worksheets/sheet1.xml': ('B2', '12'), 'xl/worksheets/sheet2.xml': ('A2', '15')}
    with ZipFile(BytesIO(buffer.getvalue())) as source, ZipFile(path, 'w') as target:
        for info in source.infolist():
            content = source.read(info)
            if info.filename in caches:
                address, value = caches[info.filename]
                sheet = ElementTree.fromstring(content)
                cell = sheet.find(f'.//{SHEET_NS}c[@r="{address}"]')
                assert cell is not None
                formula = cell.find(f'{SHEET_NS}f')
                assert formula is not None
                if flattened:
                    cell.remove(formula)
                cached = cell.find(f'{SHEET_NS}v')
                if cached is None:
                    cached = ElementTree.SubElement(cell, f'{SHEET_NS}v')
                cached.text = value
                content = ElementTree.tostring(sheet, encoding='utf-8', xml_declaration=True)
            target.writestr(info, content)


def test_published_xlsx_retains_live_formulas_and_expected_calculated_values(tmp_path: Path) -> None:
    published = tmp_path / 'published.xlsx'
    _published_workbook(published)

    observed_values = judge_xlsx_formulas(
        published,
        expected_formulas=FORMULAS,
        expected_values=EXPECTED_VALUES,
    )

    assert observed_values == {('Main', 'B2'): 12, ('Second', 'A2'): 15}


def test_flattened_xlsx_fails_formula_check_despite_correct_displayed_values(tmp_path: Path) -> None:
    published = tmp_path / 'published.xlsx'
    _published_workbook(published, flattened=True)
    displayed = load_workbook(published, data_only=True, read_only=True)
    try:
        assert displayed['Main']['B2'].value == 12
        assert displayed['Second']['A2'].value == 15
    finally:
        displayed.close()

    with pytest.raises(DocumentCheckError) as failure:
        judge_xlsx_formulas(
            published,
            expected_formulas=FORMULAS,
            expected_values=EXPECTED_VALUES,
        )

    assert failure.value.format == 'XLSX'
    assert failure.value.check == 'formulas'
    assert 'XLSX' in str(failure.value)
    assert 'formula' in str(failure.value).lower()


def _image() -> BytesIO:
    image = Image.new('RGB', (16, 16), (35, 100, 170))
    for y in range(16):
        for x in range(8):
            image.putpixel((x, y), (220, 40, 50))
    stream = BytesIO()
    image.save(stream, format='PNG')
    stream.seek(0)
    return stream


def _published_docx(path: Path, *, chinese=True) -> None:
    document = Document()
    document.add_paragraph('OCU_DOCX_EDIT Editable introduction')
    document.add_heading('Office verification', level=1)
    document.add_heading('Content preservation', level=2)
    if chinese:
        document.add_paragraph('中文内容保持完整。')
    table = document.add_table(rows=2, cols=2)
    for row, cells in zip(table.rows, (('Item', 'Count'), ('中文数据', '3')), strict=True):
        for cell, value in zip(row.cells, cells, strict=True):
            cell.text = value
    document.add_picture(_image())
    document.save(path)


def test_docx_content_judge_accepts_preserved_content_and_rejects_lost_chinese(tmp_path):
    valid = tmp_path / 'valid.docx'
    damaged = tmp_path / 'damaged.docx'
    _published_docx(valid)
    _published_docx(damaged, chinese=False)
    assert judge_docx(valid) == {
        'format': 'DOCX',
        'checks': ['edit', 'headings', 'table', 'image', 'Chinese text'],
    }
    with pytest.raises(DocumentCheckError) as failure:
        judge_docx(damaged)
    assert (failure.value.format, failure.value.check) == ('DOCX', 'Chinese text')


def test_docx_content_judge_rejects_collaboration_name_in_edited_introduction(tmp_path):
    published = tmp_path / 'contaminated.docx'
    _published_docx(published)
    document = Document(published)
    document.paragraphs[0].text = 'OCU_DOCX_EDIT OCU office verificationEditable introduction'
    document.save(published)

    with pytest.raises(DocumentCheckError) as failure:
        judge_docx(published)

    assert (failure.value.format, failure.value.check) == ('DOCX', 'edit')


def _published_docx_choice(path: Path, *, fallback_picture=False) -> None:
    _published_docx(path)
    document = Document(path)
    drawing = document.inline_shapes[0]._inline.getparent()
    run = drawing.getparent()
    alternate = etree.Element(
        f'{{{MC_NS}}}AlternateContent',
        nsmap={
            'mc': MC_NS,
            'wpg': 'http://schemas.microsoft.com/office/word/2010/wordprocessingGroup',
            'unsupported': 'urn:unsupported-drawing',
        },
    )
    run.replace(drawing, alternate)
    choice = etree.SubElement(alternate, f'{{{alternate.nsmap["mc"]}}}Choice', Requires='wpg')
    choice.append(drawing)
    fallback = etree.SubElement(alternate, f'{{{alternate.nsmap["mc"]}}}Fallback')
    if fallback_picture:
        fallback.append(deepcopy(drawing))
    document.save(path)


@pytest.mark.parametrize('alternative', ['empty fallback', 'picture fallback', 'later choice'])
def test_docx_content_judge_accepts_selected_choice_image_without_counting_alternatives(tmp_path, alternative):
    published = tmp_path / 'published.docx'
    _published_docx_choice(published, fallback_picture=alternative == 'picture fallback')
    if alternative == 'later choice':
        document = Document(published)
        alternate = next(document.element.iter(f'{{{MC_NS}}}AlternateContent'))
        alternate.insert(1, deepcopy(alternate[0]))
        document.save(published)
    before = published.read_bytes()

    assert judge_docx(published) == {
        'format': 'DOCX',
        'checks': ['edit', 'headings', 'table', 'image', 'Chinese text'],
    }
    assert published.read_bytes() == before


@pytest.mark.parametrize(
    'damage',
    [
        'missing picture',
        'dangling relationship',
        'wrong relationship',
        'changed pixels',
        'changed dimensions',
        'unreadable media',
        'inactive fallback',
        'unsupported choice',
    ],
)
def test_docx_content_judge_rejects_missing_or_damaged_selected_choice_image(tmp_path, damage):
    published = tmp_path / 'damaged.docx'
    _published_docx_choice(published, fallback_picture=damage == 'inactive fallback')
    document = Document(published)
    alternate = next(document.element.iter(f'{{{MC_NS}}}AlternateContent'))
    choice = alternate[0]
    drawing = choice[0]
    blip = document.element.xpath('//a:blip')[0]
    relationship = blip.get(qn('r:embed'))
    if damage in ('missing picture', 'inactive fallback'):
        choice.remove(drawing)
    elif damage == 'dangling relationship':
        blip.set(qn('r:embed'), 'rIdMissing')
    elif damage == 'wrong relationship':
        blip.set(
            qn('r:embed'),
            next(key for key, part in document.part.related_parts.items() if part.content_type.endswith('styles+xml')),
        )
    elif damage == 'unsupported choice':
        choice.set('Requires', 'unsupported')
    else:
        image_part = document.part.related_parts[relationship]
        if damage == 'unreadable media':
            image_part._blob = b'not an image'
        else:
            image = Image.new('RGB', (16, 16) if damage == 'changed pixels' else (8, 16), (35, 100, 170))
            stream = BytesIO()
            image.save(stream, format='PNG')
            image_part._blob = stream.getvalue()
    document.save(published)
    before = published.read_bytes()

    with pytest.raises(DocumentCheckError) as failure:
        judge_docx(published)

    assert (failure.value.format, failure.value.check) == ('DOCX', 'image')
    assert published.read_bytes() == before


def _published_pptx(path: Path, *, reversed_slides=False) -> None:
    deck = Presentation()
    for title in ('Slide One OCU_PPTX_EDIT', 'Slide Two', 'Slide Three'):
        deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = title
    rectangle = deck.slides[0].shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(1),
        Inches(5),
        Inches(3),
        Inches(1),
    )
    rectangle.text = 'Preserved shape'
    deck.slides[1].shapes.add_picture(_image(), Inches(1), Inches(5))
    if reversed_slides:
        order = deck.slides._sldIdLst
        order.insert(1, order[2])
    deck.save(path)


def test_pptx_content_judge_accepts_preserved_slides_and_rejects_reordering(tmp_path):
    valid = tmp_path / 'valid.pptx'
    damaged = tmp_path / 'damaged.pptx'
    _published_pptx(valid)
    _published_pptx(damaged, reversed_slides=True)
    assert judge_pptx(valid) == {
        'format': 'PPTX',
        'checks': ['edit', 'slide count', 'slide order', 'shapes', 'image'],
    }
    with pytest.raises(DocumentCheckError) as failure:
        judge_pptx(damaged)
    assert (failure.value.format, failure.value.check) == ('PPTX', 'slide order')


def test_xlsx_content_judge_accepts_cross_sheet_values_and_rejects_stale_cache(tmp_path):
    published = tmp_path / 'published.xlsx'
    _published_workbook(published)
    assert judge_xlsx(published)['values'] == {'Main!B2': 12, 'Second!A2': 15}
    changed = tmp_path / 'stale.xlsx'
    with ZipFile(published) as source, ZipFile(changed, 'w') as target:
        for info in source.infolist():
            content = source.read(info)
            if info.filename == 'xl/worksheets/sheet2.xml':
                tree = ElementTree.fromstring(content)
                tree.find(f'.//{SHEET_NS}c[@r="A2"]/{SHEET_NS}v').text = '9'
                content = ElementTree.tostring(tree)
            target.writestr(info, content)
    with pytest.raises(DocumentCheckError) as failure:
        judge_xlsx(changed)
    assert (failure.value.format, failure.value.check) == ('XLSX', 'values')
